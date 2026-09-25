"""
documents.py — NOVA's notebook.

take_note()     append a timestamped note to logs/notes/<file>.txt
list_notes()    show what's in the notebook
read_note()     read the most recent note back
"""

import os
import time

import config


def _note_path(title: str) -> str:
    safe = "".join(c for c in title if c.isalnum() or c in " -_").strip() or "notes"
    if not safe.endswith(".txt"):
        safe += ".txt"
    return os.path.join(config.NOTES_DIR, safe)


def take_note(title: str, body: str = ""):
    path = _note_path(title or "quick note")
    stamp = time.strftime("%Y-%m-%d %H:%M")
    with open(path, "a", encoding="utf-8") as f:
        f.write(f"[{stamp}] {body or title}\n")
    return f"Noted. It's saved in {os.path.basename(path)}."


def list_notes():
    files = sorted(os.listdir(config.NOTES_DIR))
    if not files:
        return "Your notebook is empty."
    names = ", ".join(os.path.splitext(f)[0] for f in files[:8])
    return f"You have {len(files)} note(s): {names}."


def read_note(title: str = ""):
    if title:
        path = _note_path(title)
        if not os.path.exists(path):
            return f"I can't find a note called {title}."
    else:
        files = sorted(
            (os.path.join(config.NOTES_DIR, f) for f in os.listdir(config.NOTES_DIR)),
            key=os.path.getmtime,
        )
        if not files:
            return "Your notebook is empty."
        path = files[-1]

    with open(path, "r", encoding="utf-8") as f:
        lines = [l.strip() for l in f if l.strip()]
    if not lines:
        return "That note is empty."
    return f"Latest entry: {lines[-1]}"


def write_document(title: str, body: str, fmt: str = "txt"):
    """
    Create a real document in logs/notes as .txt, .docx (python-docx)
    or .pdf (fpdf2). Degrades gracefully if a library is missing.
    """
    safe = "".join(c for c in title if c.isalnum() or c in " -_").strip() or "document"
    path = os.path.join(config.NOTES_DIR, safe + "." + fmt)
    try:
        if fmt == "txt":
            with open(path, "w", encoding="utf-8") as f:
                f.write(body)
        elif fmt == "docx":
            import docx

            d = docx.Document()
            d.add_heading(title, level=1)
            for para in body.splitlines():
                if para.strip():
                    d.add_paragraph(para)
            d.save(path)
        elif fmt == "pdf":
            from fpdf import FPDF

            pdf = FPDF()
            pdf.add_page()
            pdf.set_font("Helvetica", size=12)
            pdf.multi_cell(0, 8, title)
            pdf.multi_cell(0, 8, "")
            pdf.multi_cell(0, 8, body)
            pdf.output(path)
        else:
            return f"I don't know the {fmt} format yet, sir."
        return f"Done. The document is saved as {os.path.basename(path)}."
    except ImportError:
        hint = "python-docx" if fmt == "docx" else "fpdf2"
        return f"Creating {fmt} files needs one extra package — pip install {hint}."
    except Exception as e:
        return f"The document didn't save, sir: {e}"
